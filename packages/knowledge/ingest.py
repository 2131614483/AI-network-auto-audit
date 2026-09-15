"""Deterministic local-file knowledge ingestion (Phase 2 boundary)."""

from __future__ import annotations

import csv
import hashlib
import os
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

import psycopg2
from psycopg2.extras import Json

from packages.knowledge.local_extractors import (
    MEDIA_SUFFIXES,
    MINERU_SUFFIXES,
    deferred_extraction_for,
)

SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".docx", ".xlsx"} | MINERU_SUFFIXES | MEDIA_SUFFIXES

_MIME_TYPES = {
    ".csv": "text/csv",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pdf": "application/pdf",
}


@dataclass(frozen=True, slots=True)
class IngestedFile:
    source_uri: str
    sha256: str
    byte_size: int
    chunks: int
    status: str = "accepted"
    extractor_key: str | None = None


@dataclass(frozen=True, slots=True)
class IngestResult:
    batch_id: str | None
    scanned: int
    accepted: int
    skipped: int
    deferred: int
    files: tuple[IngestedFile, ...]


@dataclass(frozen=True, slots=True)
class _PreparedFile:
    path: Path
    item: IngestedFile
    title: str
    chunks: tuple[str, ...]
    status: str
    extractor_key: str | None
    error_detail: str | None


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _parse_text(data: bytes, fallback_title: str) -> tuple[str, str]:
    """Decode text and remove optional YAML front matter from indexed body."""
    text = data.decode("utf-8", errors="replace").replace("\r\n", "\n").replace("\r", "\n")
    title = fallback_title
    if text.startswith("---\n"):
        marker = text.find("\n---\n", 4)
        if marker >= 0:
            for line in text[4:marker].splitlines():
                key, separator, value = line.partition(":")
                if separator and key.strip().lower() == "title" and value.strip():
                    title = value.strip().strip('"\'')
                    break
            text = text[marker + len("\n---\n") :]
    return title, text.strip()


def _parse_csv(data: bytes, fallback_title: str) -> tuple[str, str]:
    """Turn a local CSV into a deterministic, readable table transcript."""
    text = data.decode("utf-8-sig", errors="replace")
    rows = csv.reader(StringIO(text))
    rendered = [" | ".join(cell.strip() for cell in row) for row in rows]
    return fallback_title, "\n".join(line for line in rendered if line).strip()


def _parse_docx(data: bytes, fallback_title: str) -> tuple[str, str]:
    """Extract paragraphs from DOCX without executing embedded content."""
    try:
        from xml.etree import ElementTree

        with ZipFile(BytesIO(data)) as archive:
            document = archive.read("word/document.xml")
        root = ElementTree.fromstring(document)
    except (BadZipFile, KeyError, ValueError) as exc:
        raise ValueError("invalid DOCX document") from exc
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs = []
    for paragraph in root.iter(f"{namespace}p"):
        content = "".join(node.text or "" for node in paragraph.iter(f"{namespace}t")).strip()
        if content:
            paragraphs.append(content)
    return fallback_title, "\n\n".join(paragraphs)


def _parse_xlsx(data: bytes, fallback_title: str) -> tuple[str, str]:
    """Extract cell values only; formulas and macros are never executed."""
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, ValueError) as exc:
        raise ValueError("invalid XLSX workbook") from exc
    try:
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            lines.append(f"# Sheet: {worksheet.title}")
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value).strip() for value in row]
                if any(values):
                    lines.append(" | ".join(values))
        return fallback_title, "\n".join(lines)
    finally:
        workbook.close()


def _parse_document(path: Path, data: bytes) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        return _parse_text(data, path.stem)
    if suffix == ".csv":
        return _parse_csv(data, path.stem)
    if suffix == ".docx":
        return _parse_docx(data, path.stem)
    if suffix == ".xlsx":
        return _parse_xlsx(data, path.stem)
    raise ValueError(f"unsupported knowledge file type: {suffix}")


def chunk_text(text: str, *, max_chars: int = 1600) -> list[str]:
    """Split text on paragraph boundaries without losing content."""
    if max_chars < 128:
        raise ValueError("max_chars must be at least 128")
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [part.strip() for part in normalized.split("\n\n") if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(paragraph[i : i + max_chars] for i in range(0, len(paragraph), max_chars))
        elif not current:
            current = paragraph
        elif len(current) + 2 + len(paragraph) <= max_chars:
            current = f"{current}\n\n{paragraph}"
        else:
            chunks.append(current)
            current = paragraph
    if current:
        chunks.append(current)
    return chunks


def scan_directory(root: Path) -> list[Path]:
    if not root.exists() or not root.is_dir():
        raise ValueError(f"knowledge drop directory does not exist: {root}")
    return sorted(
        path for path in root.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )


def _safe_relative(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError as exc:
        raise ValueError("file escapes knowledge drop root") from exc


def _prepare(root: Path, paths: list[Path], max_chars: int) -> list[_PreparedFile]:
    namespace = _sha256(str(root).encode())[:16]
    result: list[_PreparedFile] = []
    for path in paths:
        data = path.read_bytes()
        deferred = deferred_extraction_for(path)
        if deferred is None:
            title, body = _parse_document(path, data)
            chunks = tuple(chunk_text(body, max_chars=max_chars))
            status, extractor_key, error_detail = "accepted", None, None
        else:
            title, chunks = path.stem, ()
            status, extractor_key, error_detail = "waiting_extractor", deferred.extractor_key, deferred.reason
        result.append(
            _PreparedFile(
                path,
                IngestedFile(
                    f"drop://{namespace}/{_safe_relative(path, root)}",
                    _sha256(data),
                    len(data),
                    len(chunks),
                    status,
                    extractor_key,
                ),
                title,
                chunks,
                status,
                extractor_key,
                error_detail,
            )
        )
    return result


def _artifact_id(cur: Any, tenant_id: Any, item: IngestedFile) -> Any:
    cur.execute(
        "SELECT id FROM artifact.artifacts WHERE tenant_id=%s AND blob_sha256=%s ORDER BY created_at LIMIT 1",
        (tenant_id, item.sha256),
    )
    existing = cur.fetchone()
    if existing:
        return existing[0]
    cur.execute(
        """INSERT INTO artifact.artifacts
        (tenant_id,blob_sha256,artifact_type,classification,metadata)
        VALUES(%s,%s,'source_document','internal',%s) RETURNING id""",
        (tenant_id, item.sha256, Json({"ingest": "phase2"})),
    )
    row = cur.fetchone()
    if row is None:
        raise RuntimeError("artifact insert returned no id")
    return row[0]


def ingest_directory(
    root: str | Path,
    *,
    database_url: str,
    tenant_slug: str = "local-dev",
    idempotency_key: str | None = None,
    max_chars: int = 1600,
    dry_run: bool = False,
) -> IngestResult:
    """Atomically stage files; identical content is skipped on retries."""
    root_path = Path(root).resolve()
    paths = scan_directory(root_path)
    prepared = _prepare(root_path, paths, max_chars)
    files = tuple(item.item for item in prepared)
    if dry_run or not files:
        return IngestResult(
            None,
            len(paths),
            sum(item.status == "accepted" for item in prepared),
            0,
            sum(item.status == "waiting_extractor" for item in prepared),
            files,
        )

    key = idempotency_key or _sha256("\n".join(f.source_uri + f.sha256 for f in files).encode())
    connection = psycopg2.connect(database_url)
    try:
        with connection:
            with connection.cursor() as cur:
                cur.execute("SELECT id FROM iam.tenants WHERE slug=%s", (tenant_slug,))
                tenant_row = cur.fetchone()
                if not tenant_row:
                    raise ValueError(f"tenant not found: {tenant_slug}")
                tenant_id = tenant_row[0]
                cur.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant_id),))
                cur.execute(
                    """INSERT INTO knowledge.ingest_batches
                    (tenant_id,source_uri,idempotency_key,status,scanned_count,started_at)
                    VALUES(%s,%s,%s,'running',%s,now())
                    ON CONFLICT (tenant_id,idempotency_key) DO UPDATE SET
                      source_uri=EXCLUDED.source_uri,status='running',scanned_count=EXCLUDED.scanned_count,
                      accepted_count=0,skipped_count=0,failed_count=0,error_detail=NULL,
                      started_at=now(),finished_at=NULL RETURNING id""",
                    (tenant_id, str(root_path), key, len(files)),
                )
                batch = cur.fetchone()
                if batch is None:
                    raise RuntimeError("ingest batch upsert returned no id")
                batch_uuid = batch[0]
                accepted = 0
                skipped = 0
                deferred_count = 0
                for item in prepared:
                    relative = _safe_relative(item.path, root_path)
                    cur.execute(
                        """INSERT INTO knowledge.ingest_files
                        (tenant_id,batch_id,relative_path,source_uri,sha256,size_bytes,mime_type,status)
                        VALUES(%s,%s,%s,%s,%s,%s,%s,%s)
                        ON CONFLICT(batch_id,relative_path) DO UPDATE SET source_uri=EXCLUDED.source_uri,
                        sha256=EXCLUDED.sha256,size_bytes=EXCLUDED.size_bytes,status=EXCLUDED.status,
                        error_code=NULL,error_detail=NULL RETURNING id""",
                        (tenant_id, batch_uuid, relative, item.item.source_uri, item.item.sha256, item.item.byte_size,
                         _MIME_TYPES.get(item.path.suffix.lower(), "text/plain"), item.status),
                    )
                    cur.fetchone()
                    cur.execute(
                        """INSERT INTO artifact.blobs(sha256,byte_size,media_type,storage_uri)
                        VALUES(%s,%s,%s,%s) ON CONFLICT DO NOTHING""",
                        (item.item.sha256, item.item.byte_size,
                         _MIME_TYPES.get(item.path.suffix.lower(), "text/plain"), f"file://{item.path}"),
                    )
                    artifact_id = _artifact_id(cur, tenant_id, item.item)
                    if item.status == "waiting_extractor":
                        cur.execute(
                            """UPDATE knowledge.ingest_files SET error_code=%s,error_detail=%s,source_artifact_id=%s,
                            finished_at=now() WHERE batch_id=%s AND relative_path=%s""",
                            (item.extractor_key, item.error_detail, artifact_id, batch_uuid, relative),
                        )
                        deferred_count += 1
                        continue
                    cur.execute(
                        """SELECT id,content_sha256,current_version FROM semantic.documents
                        WHERE tenant_id=%s AND source_uri=%s FOR UPDATE""",
                        (tenant_id, item.item.source_uri),
                    )
                    existing = cur.fetchone()
                    if existing and existing[1] == item.item.sha256:
                        document_id, document_version = existing[0], existing[2]
                        skipped += 1
                        status = "skipped"
                    elif existing:
                        document_id, document_version = existing[0], existing[2] + 1
                        cur.execute(
                            """UPDATE semantic.documents SET title=%s,status='staged',current_version=%s,
                            content_sha256=%s,source_artifact_id=%s,updated_at=now() WHERE id=%s""",
                            (item.title, document_version, item.item.sha256, artifact_id, document_id),
                        )
                        for ordinal, chunk in enumerate(item.chunks):
                            cur.execute(
                                """INSERT INTO semantic.chunks
                                (tenant_id,document_id,document_version,ordinal,content,token_count,metadata)
                                VALUES(%s,%s,%s,%s,%s,%s,%s)""",
                                (tenant_id, document_id, document_version, ordinal, chunk, len(chunk), Json({"source": "phase2"})),
                            )
                        accepted += 1
                        status = "accepted"
                    else:
                        document_version = 1
                        cur.execute(
                            """INSERT INTO semantic.documents
                            (tenant_id,source_uri,title,status,current_version,content_sha256,source_artifact_id)
                            VALUES(%s,%s,%s,'staged',1,%s,%s) RETURNING id""",
                            (tenant_id, item.item.source_uri, item.title, item.item.sha256, artifact_id),
                        )
                        inserted_document = cur.fetchone()
                        if inserted_document is None:
                            raise RuntimeError("document insert returned no id")
                        document_id = inserted_document[0]
                        for ordinal, chunk in enumerate(item.chunks):
                            cur.execute(
                                """INSERT INTO semantic.chunks
                                (tenant_id,document_id,document_version,ordinal,content,token_count,metadata)
                                VALUES(%s,%s,%s,%s,%s,%s,%s)""",
                                (tenant_id, document_id, document_version, ordinal, chunk, len(chunk), Json({"source": "phase2"})),
                            )
                        accepted += 1
                        status = "accepted"
                    cur.execute(
                        """UPDATE knowledge.ingest_files SET status=%s,source_artifact_id=%s,
                        document_id=%s,finished_at=now() WHERE batch_id=%s AND relative_path=%s""",
                        (status, artifact_id, document_id, batch_uuid, relative),
                    )
                cur.execute(
                    """UPDATE knowledge.ingest_batches SET status='completed',accepted_count=%s,
                    skipped_count=%s,deferred_count=%s,finished_at=now() WHERE id=%s""",
                    (accepted, skipped, deferred_count, batch_uuid),
                )
        return IngestResult(str(batch_uuid), len(paths), accepted, skipped, deferred_count, files)
    finally:
        connection.close()


def default_drop_root() -> Path:
    return Path(os.getenv("KNOWLEDGE_DROP_ROOT", ".data/drop"))
